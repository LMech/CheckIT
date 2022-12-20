from PyPDF2 import PdfReader
from openpyxl import load_workbook
from docx import Document
from PIL import Image


def pdfCorrupted(path: str) -> bool:
    with open(path, "rb") as f:
        try:
            PdfReader(f)
            return False
        except:
            return True


def exclCorrupted(path: str) -> bool:
    try:
        load_workbook(
            path,
        )
        return False
    except:
        return True


def docxCorrupted(path: str) -> bool:
    try:
        Document(path)
        return False
    except:
        return True


def jpegCorrupted(path: str) -> bool:
    try:
        im = Image.open(path)
        im.verify()
        return False
    except:
        return True


def is_corrupted(path: str):
    if path.endswith("docx"):
        return docxCorrupted(path)
    elif path.endswith("xlsx"):
        return exclCorrupted(path)
    elif path.endswith(".pdf"):
        return pdfCorrupted(path)
    elif path.endswith(".jpg") or path.endswith(".jpeg"):
        return jpegCorrupted(path)
